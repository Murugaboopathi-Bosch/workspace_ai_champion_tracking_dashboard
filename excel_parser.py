"""
excel_parser.py
Parses the Gen AI Champions Initiative Excel sheet into structured data.

Handles the real-world messy layout:
  Region 1: header block (Initiative Name / Owner / Supported By / Champions / Network)
  Region 2: Topic / Sub topics / Progress / Lead / Key Points / Start Date / End Date /
            Duration / Docupedia Link table (with topic-group rollups)
  Region 3: freeform notes after the table

Optional second sheet (any of: "Activity Sessions", "Activity_Sessions", "ActivitySessions",
"Sessions", "Tech_Talk_Sessions") provides child records (e.g. individual AI Tech Talk
sessions) linked to a Sub Topic / Activity by name match.
"""

import re
from dataclasses import dataclass, field
from datetime import datetime
from openpyxl import load_workbook


@dataclass
class ChampionMember:
    name: str
    department: str | None = None
    group: str = "AI Champions"  # or "AI Champions Network"


@dataclass
class ActivitySession:
    """Represents a child session/record for an activity (e.g., individual tech talks)."""
    session_name: str
    activity_name: str  # Links back to the parent subtopic by name
    status: str | None = None
    start_date: str | None = None
    end_date: str | None = None
    duration: str | None = None
    timeline: str | None = None
    notes: str | None = None


@dataclass
class Subtopic:
    name: str
    progress_pct: float | None = None
    lead: str | None = None
    key_points: str | None = None
    start_date: str | None = None
    end_date: str | None = None
    duration: str | None = None
    docupedia_link: str | None = None
    sessions: list[ActivitySession] = field(default_factory=list)


@dataclass
class TopicGroup:
    name: str
    overall_progress_pct: float | None = None
    subtopics: list[Subtopic] = field(default_factory=list)


@dataclass
class DashboardData:
    initiative_name: str
    owner: str | None
    supported_by: str | None
    champions: list[ChampionMember]
    topics: list[TopicGroup]
    notes: list[str]
    last_synced: datetime
    source_file: str
    activity_sessions: list[ActivitySession] = field(default_factory=list)


NAME_DEPT_RE = re.compile(r"^(.*?)\s*\(([^)]+)\)\s*$")

# Any of these sheet names will be recognized for the child-sessions table,
# case-insensitively, so it doesn't matter which exact naming convention was used.
SESSIONS_SHEET_CANDIDATES = [
    "Activity Sessions", "Activity_Sessions", "ActivitySessions",
    "Sessions", "Tech_Talk_Sessions", "Tech Talk Sessions",
]


def _split_name_dept(raw: str) -> tuple[str, str | None]:
    if not raw:
        return raw, None
    m = NAME_DEPT_RE.match(raw.strip())
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return raw.strip(), None


def _is_blank_row(row) -> bool:
    return all(c is None or str(c).strip() == "" for c in row)


def _pct(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    return round(v * 100, 1) if v <= 1 else round(v, 1)


def _format_date(date_val) -> str | None:
    """Convert a date value (real Excel date or text) to a display string."""
    if date_val is None or str(date_val).strip() == "":
        return None
    if isinstance(date_val, datetime):
        return date_val.strftime('%Y-%m-%d')
    return str(date_val).strip()


def _find_sheet_flexible(wb, candidates: list[str]):
    """Case-insensitive, whitespace/underscore-insensitive sheet name lookup."""
    def normalize(s):
        return s.lower().replace("_", " ").strip()
    norm_map = {normalize(name): name for name in wb.sheetnames}
    for cand in candidates:
        n = normalize(cand)
        if n in norm_map:
            return wb[norm_map[n]]
    return None


def _extract_hyperlink_map(ws, col_idx_0based: int) -> dict:
    """Map {excel_row_number: url} for real Excel hyperlinks (Insert > Link) in a column.
    Needed because values_only=True (used elsewhere for speed) only returns the
    DISPLAYED text of a hyperlinked cell, not the actual URL it points to -- if
    someone used Insert Link with friendly text like "View Docupedia Page", the
    plain cell value is that text, not a usable link."""
    mapping = {}
    for row in ws.iter_rows():
        if len(row) > col_idx_0based:
            cell = row[col_idx_0based]
            if cell.hyperlink and cell.hyperlink.target:
                mapping[cell.row] = cell.hyperlink.target
    return mapping


def parse_excel(filepath: str) -> DashboardData:
    wb = load_workbook(filepath, data_only=True)
    ws = wb[wb.sheetnames[0]]  # always take the first sheet, don't hardcode a name
    rows = list(ws.iter_rows(values_only=True))

    DOCUPEDIA_COL_IDX = 8  # column I (0-based index 8) -- Docupedia Link
    docupedia_link_map = _extract_hyperlink_map(ws, DOCUPEDIA_COL_IDX)

    # ---- Region 1: header block ----
    initiative_name, owner, supported_by = None, None, None
    champions: list[ChampionMember] = []

    header_end_idx = 0
    for i, row in enumerate(rows[:15]):
        if row and str(row[0]).strip().lower() == "topic" and len(row) > 1 and str(row[1] or "").strip().lower().startswith("sub"):
            header_end_idx = i
            break
        if i == 0:
            continue  # this is the label row ("Initiative Name | Initiative Owner | ..."), skip it
        if row:
            if len(row) > 0 and row[0] and initiative_name is None:
                initiative_name = row[0]
            if len(row) > 1 and row[1] and owner is None:
                owner = row[1]
            if len(row) > 2 and row[2] and supported_by is None:
                supported_by = row[2]
            if len(row) > 3 and row[3]:
                n, d = _split_name_dept(str(row[3]))
                champions.append(ChampionMember(name=n, department=d, group="AI Champions"))
            if len(row) > 4 and row[4]:
                n, d = _split_name_dept(str(row[4]))
                champions.append(ChampionMember(name=n, department=d, group="AI Champions Network"))

    # ---- Region 2: topics table ----
    # Column layout: A Topic | B Sub topics | C Progress | D Lead | E Key Points |
    #                F Start Date | G End Date | H Duration | I Docupedia Link
    # NOTE: there is deliberately NO placeholder/stray column between Key Points and
    # Start Date -- an earlier version had a dead "stray_val" read here that silently
    # shifted every field after it by one column. Don't reintroduce that.
    topics: list[TopicGroup] = []
    current_group: TopicGroup | None = None
    notes: list[str] = []
    in_notes_section = False
    blank_run = 0

    for abs_idx, row in enumerate(rows[header_end_idx + 1:], start=header_end_idx + 1):
        excel_row_num = abs_idx + 1  # openpyxl rows are 1-based

        if _is_blank_row(row):
            blank_run += 1
            if blank_run >= 2 and current_group is not None:
                in_notes_section = True
            continue
        blank_run = 0

        topic_val = row[0] if len(row) > 0 else None
        subtopic_val = row[1] if len(row) > 1 else None
        progress_val = row[2] if len(row) > 2 else None
        lead_val = row[3] if len(row) > 3 else None
        keypoints_val = row[4] if len(row) > 4 else None
        start_date_val = row[5] if len(row) > 5 else None
        end_date_val = row[6] if len(row) > 6 else None
        duration_val = row[7] if len(row) > 7 else None
        docupedia_val = row[8] if len(row) > 8 else None

        is_note_row = (
            subtopic_val is not None
            and progress_val is None
            and lead_val is None
            and keypoints_val is None
            and topic_val is None
            and start_date_val is None
            and end_date_val is None
            and duration_val is None
            and docupedia_val is None
        )
        if in_notes_section or is_note_row:
            if subtopic_val:
                notes.append(str(subtopic_val).strip())
            continue

        if topic_val and not subtopic_val:
            current_group = TopicGroup(
                name=str(topic_val).strip(),
                overall_progress_pct=_pct(progress_val),
            )
            topics.append(current_group)
            continue

        if subtopic_val:
            if current_group is None:
                current_group = TopicGroup(name="General", overall_progress_pct=None)
                topics.append(current_group)

            # Prefer the real hyperlink target (set via Excel's Insert Link) over the
            # plain cell text, since the displayed text is often just a label like
            # "View Docupedia Page" rather than a usable URL.
            link_final = docupedia_link_map.get(excel_row_num) or (str(docupedia_val).strip() if docupedia_val else None)

            current_group.subtopics.append(Subtopic(
                name=str(subtopic_val).strip(),
                progress_pct=_pct(progress_val),
                lead=str(lead_val).strip() if lead_val else None,
                key_points=str(keypoints_val).strip() if keypoints_val else None,
                start_date=_format_date(start_date_val),
                end_date=_format_date(end_date_val),
                duration=str(duration_val).strip() if duration_val else None,
                docupedia_link=link_final,
            ))

    activity_sessions = _parse_activity_sessions(wb)
    _link_sessions_to_activities(topics, activity_sessions)

    return DashboardData(
        initiative_name=str(initiative_name) if initiative_name else "Untitled Initiative",
        owner=str(owner) if owner else None,
        supported_by=str(supported_by) if supported_by else None,
        champions=champions,
        topics=topics,
        notes=notes,
        last_synced=datetime.now(),
        source_file=filepath,
        activity_sessions=activity_sessions,
    )


# ---------------------------------------------------------------------------
# Activity Sessions parsing (flexible header matching)
# ---------------------------------------------------------------------------

FIELD_HEADER_ALIASES = {
    "activity_name": ["activity name", "parent activity", "sub topic", "subtopic", "sub-topic", "activity"],
    "session_name":  ["session name", "session title", "title", "session"],
    "status":        ["status"],
    "start_date":    ["start date", "start"],
    "end_date":      ["end date", "end"],
    "duration":      ["duration", "elapsed"],
    "timeline":      ["timeline"],
    "notes":         ["notes", "note", "remarks"],
}


def _match_field(header_cell: str) -> str | None:
    h = header_cell.strip().lower()
    for field_name, aliases in FIELD_HEADER_ALIASES.items():
        for alias in aliases:
            if alias in h:
                return field_name
    return None


def _parse_activity_sessions(workbook) -> list[ActivitySession]:
    """Parse the optional child-sessions sheet, if present. Returns [] gracefully if
    it doesn't exist or can't be matched -- older workbooks without this sheet still work."""
    ws = _find_sheet_flexible(workbook, SESSIONS_SHEET_CANDIDATES)
    if ws is None:
        return []

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_idx = None
    col_indices = {}
    for i, row in enumerate(rows[:10]):
        if not row:
            continue
        candidate_indices = {}
        for col_i, cell in enumerate(row):
            if cell is None:
                continue
            matched = _match_field(str(cell))
            if matched and matched not in candidate_indices:
                candidate_indices[matched] = col_i
        if "activity_name" in candidate_indices and "session_name" in candidate_indices:
            header_idx = i
            col_indices = candidate_indices
            break

    if header_idx is None:
        return []

    sessions: list[ActivitySession] = []
    for row in rows[header_idx + 1:]:
        if _is_blank_row(row):
            continue

        def get(field_name):
            idx = col_indices.get(field_name)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        activity_name = get("activity_name")
        session_name = get("session_name")
        if not activity_name or not session_name:
            continue

        sessions.append(ActivitySession(
            activity_name=str(activity_name).strip(),
            session_name=str(session_name).strip(),
            status=str(get("status")).strip() if get("status") else None,
            start_date=_format_date(get("start_date")),
            end_date=_format_date(get("end_date")),
            duration=str(get("duration")).strip() if get("duration") else None,
            timeline=str(get("timeline")).strip() if get("timeline") else None,
            notes=str(get("notes")).strip() if get("notes") else None,
        ))

    return sessions


def _link_sessions_to_activities(topics: list[TopicGroup], sessions: list[ActivitySession]) -> None:
    """Link activity sessions to their parent subtopics by (normalized) name match."""
    activity_map = {}
    for topic in topics:
        for subtopic in topic.subtopics:
            activity_map[subtopic.name.strip().lower()] = subtopic

    for session in sessions:
        normalized_activity = session.activity_name.strip().lower()
        if normalized_activity in activity_map:
            activity_map[normalized_activity].sessions.append(session)